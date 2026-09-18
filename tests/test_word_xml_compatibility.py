import copy
from datetime import datetime
import io
import unittest
import zipfile
from unittest.mock import patch

from docx import Document
from lxml import etree
from app.router import postgresql as report


class WordXmlCompatibilityTests(unittest.TestCase):
    def test_xml_forbidden_characters_only(self):
        invalid = ''.join(chr(i) for i in range(32) if i not in (9, 10, 13))
        invalid += '\ud800\udfff\ufffe\uffff'
        self.assertEqual(report._word_xml_text(invalid), '')
        valid = '中文★△我方\t20万元\n2026-09-16\r<A&B>"\u0085\ue000\ufffd\U00020000😀'
        self.assertEqual(report._word_xml_text(valid), valid)
        self.assertEqual(report._word_xml_text(None), '')
        self.assertEqual(report._word_xml_text(0), '0')

    def test_paragraph_header_and_cell_round_trip(self):
        doc = Document()
        text = '原文\x00\x01\x0b\x1f\ud800\uffff保留\t制表\n换行 & <文本> 😀'
        expected = '原文保留\t制表\n换行 & <文本> 😀'
        report._add_word_paragraph(doc, text)
        report._add_header_text(doc, text)
        table = doc.add_table(rows=1, cols=1)
        report._set_word_cell_text(table.cell(0, 0), text)
        output = io.BytesIO()
        doc.save(output)
        reopened = Document(io.BytesIO(output.getvalue()))
        self.assertEqual([p.text for p in reopened.paragraphs], [expected, expected])
        self.assertEqual(reopened.tables[0].cell(0, 0).text, expected)

    def test_full_report_with_invalid_pdf_producer_and_dynamic_headings(self):
        header = {'project_name': '项目\x00名称', 'purchaser': '采购\x01人',
                  'agency': '代理\x0b机构', 'units': [{'name': '单位\x02甲', 'contact': '联系\x03人'}],
                  'source_file_properties': [{'file_type': '技术标', 'file_name': '扫描-技术..pdf',
                                             'producer': '扫描工具\x00', 'title': 'A&B <标题> 😀'}]}
        payload = {'test': {'source': '原文\x00保留'}}
        before = copy.deepcopy((header, payload))
        sections = [('test', '审查\x04项', [('公司\x05甲', [['问题\x06', '说明', '原因', '文件.pdf', '1', 'fail']])])]
        with patch.object(report, '_collect_report_issue_sections', return_value=sections):
            output = report._render_result_word_report(payload, header=header, operator_name='操作\x00员')
        self.assertEqual((header, payload), before)
        with zipfile.ZipFile(io.BytesIO(output)) as archive:
            for name in archive.namelist():
                if name.endswith('.xml'):
                    etree.fromstring(archive.read(name))
        doc = Document(io.BytesIO(output))
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('审查项', text)
        self.assertIn('公司甲', text)
        self.assertEqual(doc.core_properties.author, '操作员')
        self.assertEqual(doc.core_properties.title, '项目名称 项目审查报告')
        cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
        self.assertIn('扫描工具', cells)
        self.assertIn('A&B <标题> 😀', cells)
        self.assertIn('问题', cells)

    def test_clean_input_report_is_unchanged(self):
        header = {'project_name': '正常项目', 'units': [], 'source_file_properties': [
            {'file_name': '原件.pdf', 'producer': '正常工具', 'title': '中文★ 20万元 & <标题>'}]}
        when = datetime(2026, 9, 16)
        cleaned = report._render_result_word_report({}, header=header, exported_at=when)
        with patch.object(report, '_word_xml_text', side_effect=lambda v: '' if v is None else str(v)):
            original = report._render_result_word_report({}, header=header, exported_at=when)
        with zipfile.ZipFile(io.BytesIO(cleaned)) as a, zipfile.ZipFile(io.BytesIO(original)) as b:
            self.assertEqual(a.namelist(), b.namelist())
            for name in a.namelist():
                self.assertEqual(a.read(name), b.read(name), name)
